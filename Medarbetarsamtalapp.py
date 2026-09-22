import streamlit as st
import google.generativeai as genai
from PIL import Image
import io
import docx
from docx.shared import Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import json

# --- SIDOINSTÄLLNINGAR & STYLING ---
st.set_page_config(page_title="Medarbetarsamtal AI - Rektorsassistent", layout="wide", page_icon="🏫")

st.title("🏫 Medarbetarsamtalsassistent för Rektorer (Gemini)")
st.markdown("Omvandla handskrivna anteckningar och foton till färdiga underlag för **Blå mallen** och **Orangea målmallen** via Google Gemini.")

# HÄMTA API-NYCKEL FRÅN STREAMLIT SECRETS ELLER TEXTFÄLT
API_KEY = st.secrets.get("GEMINI_API_KEY")

with st.sidebar:
    st.header("⚙️ Inställningar")
    if not API_KEY:
        API_KEY = st.text_input("Google Gemini API Key", type="password", help="Hämta en gratis nyckel från aistudio.google.com")
        
    st.subheader("🎯 Skolans Prioriterade Mål")
    school_goals = st.text_area(
        "Ange skolans/verksamhetens aktuella prioriteringar:",
        value="- Ökad måluppfyllelse i läs- och skrivförståelse.\n- Trygghet och studiero i alla årskurser.\n- Stärkt kollegialt lärande och digital kompetens.",
        height=140
    )
    
    st.info("""
    **Inbyggda lönekriterier (Norrköpings kommun):**
    - Yrkeskunnande
    - Flexibilitet/utveckling
    - Professionellt förhållningssätt
    - Samarbetsförmåga
    - Ledarskapsförmåga inom yrket
    - Engagemang/ansvar
    """)

if API_KEY:
    genai.configure(api_key=API_KEY)

# --- HJÄLPFUNKTION: SKAPA WORD-DOKUMENT ---
def create_word_docx(blue_content, orange_data):
    doc = docx.Document()
    
    doc.add_heading('Dokumentation - Medarbetarsamtal', level=1)
    doc.add_heading('1. Samtalsanteckningar (Arbetsmiljö, Hälsa, Trivsel, Ledarskap)', level=2)
    
    sections = [
        ("Arbetsmiljö och arbetsbelastning", blue_content.get("arbetsmiljo", "")),
        ("Hälsa", blue_content.get("halsa", "")),
        ("Trivsel och motivation", blue_content.get("trivsel", "")),
        ("Ledarskap", blue_content.get("ledarskap", ""))
    ]
    
    for sec_title, sec_text in sections:
        p_head = doc.add_paragraph()
        r_head = p_head.add_run(sec_title)
        r_head.bold = True
        r_head.font.size = Pt(11)
        
        p_box = doc.add_paragraph(sec_text if sec_text else "Inga särskilda anteckningar.")
        p_box.paragraph_format.space_after = Pt(12)

    doc.add_page_break()
    doc.add_heading('2. Mål och Utvecklingsområden (SMART-modell)', level=2)
    
    table = doc.add_table(rows=1, cols=3)
    table.style = 'Table Grid'
    
    hdr_cells = table.rows[0].cells
    headers = [
        "Mål/utvecklingsområde/krav i arbetet\n(utifrån verksamhetens mål och lönekriterier)",
        "Beskrivning av åtgärder och behov för att klara förväntat resultat inklusive vems som ansvarar för vilken del.",
        "Uppföljning"
    ]
    
    for i, text in enumerate(headers):
        hdr_cells[i].text = text
        shading = parse_xml(r'<w:shd {} w:fill="E87722"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shading)
        for paragraph in hdr_cells[i].paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    
    for row_data in orange_data:
        row_cells = table.add_row().cells
        row_cells[0].text = row_data.get("mal", "")
        row_cells[1].text = row_data.get("atgarder", "")
        row_cells[2].text = row_data.get("uppfoljning", "")
        
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

# --- HUVUDYTA: INPUT ---
col1, col2 = st.columns([1, 1])

with col1:
    st.subheader("1. Indata (Foto eller text)")
    input_type = st.radio("Välj typ av indata:", ["Foto på handskrivna anteckningar", "Klistra in text/anteckningar"])
    
    uploaded_image = None
    pasted_text = ""
    
    if input_type == "Foto på handskrivna anteckningar":
        uploaded_file = st.file_uploader("Ladda upp bild på anteckningar (PNG, JPG):", type=["png", "jpg", "jpeg"])
        if uploaded_file:
            uploaded_image = Image.open(uploaded_file)
            st.image(uploaded_image, caption="Uppladdad anteckning", use_column_width=True)
    else:
        pasted_text = st.text_area("Klistra in dina råanteckningar här:", height=250)

    start_btn = st.button("🚀 Omvandla till officiell dokumentation", type="primary")

# --- AI-PROCESSING & OUTPUT ---
with col2:
    st.subheader("2. Genererad dokumentation")
    
    if start_btn:
        if not API_KEY:
            st.error("⚠️ Vänligen ange en Gemini API-nyckel eller lägg till 'GEMINI_API_KEY' i Streamlit Secrets.")
        else:
            prompt = f"""
            Du är en expertassistent för rektorer i Norrköpings kommun.
            Tolka handskrivna/råa anteckningar från medarbetarsamtal och sortera dem i två format:

            1. BLÅ MALLEN (Samtalsområden):
            Skriv ren, professionell och officiell text under följande 4 avsnitt:
            - Arbetsmiljö och arbetsbelastning
            - Hälsa
            - Trivsel och motivation
            - Ledarskap

            2. ORANGEA MALLEN (Måltabell):
            Skapa mätbara SMART-mål kopplade till:
            a) Skolans prioriteringar: {school_goals}
            b) Norrköpings lönekriterier (Yrkeskunnande, Flexibilitet/utveckling, Professionellt förhållningssätt, Samarbetsförmåga, Ledarskapsförmåga inom yrket, Engagemang/ansvar).

            Du MÅSTE svara strikt i JSON-format med följande struktur:
            {{
              "blue_template": {{
                "arbetsmiljo": "...",
                "halsa": "...",
                "trivsel": "...",
                "ledarskap": "..."
              }},
              "orange_template": [
                {{
                  "mal": "...",
                  "atgarder": "...",
                  "uppfoljning": "..."
                }}
              ]
            }}
            """
            
            with st.spinner("Tolkar anteckningar och formulerar SMART-mål med Gemini..."):
                try:
                    model = genai.GenerativeModel(
                        model_name="gemini-1.5-flash",
                        generation_config={"response_mime_type": "application/json"}
                    )
                    
                    if uploaded_image:
                        response = model.generate_content([prompt, "Här är bilden på anteckningarna:", uploaded_image])
                    else:
                        full_prompt = f"{prompt}\n\nHär är anteckningarna:\n{pasted_text}"
                        response = model.generate_content(full_prompt)
                    
                    result = json.loads(response.text)
                    st.session_state['result'] = result
                    
                except Exception as e:
                    st.error(f"Ett fel uppstod vid bearbetningen: {e}")

    if 'result' in st.session_state:
        res = st.session_state['result']
        blue = res.get("blue_template", {})
        orange = res.get("orange_template", [])
        
        tab_blue, tab_orange = st.tabs(["🟦 Blå Mallen (Samtal)", "🟧 Orangea Mallen (SMART-mål)"])
        
        with tab_blue:
            st.markdown("### 🟦 Samtalsanteckningar")
            st.markdown(f"**Arbetsmiljö och arbetsbelastning:**\n{blue.get('arbetsmiljo', '')}")
            st.markdown(f"**Hälsa:**\n{blue.get('halsa', '')}")
            st.markdown(f"**Trivsel och motivation:**\n{blue.get('trivsel', '')}")
            st.markdown(f"**Ledarskap:**\n{blue.get('ledarskap', '')}")
            
        with tab_orange:
            st.markdown("### 🟧 Måltabell (SMART-mål)")
            for i, goal in enumerate(orange, 1):
                st.markdown(f"#### Mål {i}")
                st.markdown(f"**Mål/utvecklingsområde/krav:** {goal.get('mal')}")
                st.markdown(f"**Åtgärder, behov & ansvar:** {goal.get('atgarder')}")
                st.markdown(f"**Uppföljning:** {goal.get('uppfoljning')}")
                st.divider()

        docx_file = create_word_docx(blue, orange)
        st.download_button(
            label="📥 Ladda ner som Word-dokument (.docx)",
            data=docx_file,
            file_name="Medarbetarsamtal_Dokumentation.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
